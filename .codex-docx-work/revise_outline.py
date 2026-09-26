from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path(r"C:\Users\ADMIN\Documents\GitHub\KLTN\.codex-docx-work\reference.docx")
OUTPUT = Path(r"C:\Users\ADMIN\Documents\GitHub\KLTN\23520840_23520838_DeCuongKLTN_DieuChinhTienDo.docx")


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_paragraph(paragraphs, startswith, new_text):
    for paragraph in paragraphs:
        if paragraph.text.strip().startswith(startswith):
            set_paragraph_text(paragraph, new_text)
            return True
    raise RuntimeError(f"Paragraph not found: {startswith}")


doc = Document(SOURCE)
main = doc.tables[1]
content_cell = main.rows[5].cells[0]

replace_paragraph(
    content_cell.paragraphs,
    "Mục tiêu 1 - Thiết kế kiến trúc CPU 4 lõi",
    "Mục tiêu 1 - Thiết kế kiến trúc CPU 4 lõi (multi-core) tích hợp bộ nhớ cache: thiết kế kiến trúc bộ xử lý 4 lõi RV32IMA tích hợp bộ nhớ cache (L1/L2) bằng SystemVerilog, bao gồm: (i) thiết kế IP cache L1 (lệnh/dữ liệu, riêng từng lõi) và cache L2 (dùng chung); (ii) thiết kế bộ điều khiển cache với thuật toán cấp phát/thay thế set-associative và giao thức duy trì tính nhất quán bộ nhớ đệm; (iii) thiết kế MMU cho từng lõi, gồm TLB và PTW hỗ trợ dịch địa chỉ Sv32; (iv) thiết kế kiến trúc liên kết hai tầng, trong đó AHB-Lite nối các lõi/cache L1 với cache L2, còn AXI4 nối L2, DMA, DRAM và ngoại vi; và (v) thiết kế DMA phục vụ trao đổi dữ liệu khối với thiết bị lưu trữ ngoài.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Bước 1 - Khảo sát và nghiên cứu lý thuyết:",
    "- Bước 1 - Khảo sát và nghiên cứu lý thuyết: khảo sát nhu cầu ứng dụng thực tế và các hướng nghiên cứu đa lõi RISC-V; tìm hiểu RV32I/M/A, hệ thống đa lõi, cache nhiều cấp và coherence; nghiên cứu MMU/TLB với cơ chế dịch địa chỉ Sv32; đồng thời phân tích AHB-Lite và AXI4 theo vai trò khác nhau trong hệ thống. AHB-Lite được dùng ở miền truy cập từ lõi/cache L1 đến L2 vì giao thức đơn giản, phù hợp phân xử nhiều master; AXI4 được dùng ở miền hệ thống L2-DMA-DRAM-ngoại vi nhờ hỗ trợ nhiều kênh độc lập và truyền burst. Hai chuẩn được sử dụng bổ sung theo kiến trúc phân tầng, không phải hai phương án thay thế lẫn nhau.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Bước 2 - Thiết kế và hiện thực lõi đơn:",
    "- Bước 2 - Thiết kế và hiện thực nền tảng trong tháng 9-10: hiện thực lõi RV32IMA pipeline 5 tầng; hoàn thiện MMU/TLB/PTW và kiểm chứng hit, miss, refill, page walk, page fault; xây dựng AHB-Lite interconnect phía core-L2 và AXI4 system bus phía L2-DMA-DRAM/ngoại vi; sau đó tích hợp lõi, MMU, cache L1 và giao tiếp AHB-Lite thành một khối đơn lõi ổn định trước khi nhân bốn.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Bước 3 - Mở rộng thành hệ thống 4 lõi:",
    "- Bước 3 - Mở rộng thành hệ thống 4 lõi: nhân bản khối lõi đơn đã có MMU/TLB và cache L1 thành 4 lõi, đấu nối qua AHB-Lite đến cache L2 dùng chung; kết nối miền L2 với AXI4 system bus thông qua giao diện/bridge xác định rõ, rồi tích hợp coherence, DMA và bộ điều khiển SD/SPI. Kiểm chứng riêng handshake, arbitration, back-pressure và burst transfer của từng miền bus trước khi chạy các kịch bản chia sẻ dữ liệu toàn hệ thống.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Tìm hiểu nguyên lý tổ chức hệ thống đa lõi:",
    "- Tìm hiểu nguyên lý tổ chức hệ thống đa lõi: cache L1 riêng và L2 dùng chung; coherence kiểu snoop-invalidate; MMU/TLB và page table Sv32; kiến trúc bus hai tầng gồm AHB-Lite cho miền core/L1-L2 và AXI4 cho miền L2-DMA-DRAM-ngoại vi; cơ chế bridge giữa hai miền; DMA; và luồng nạp chương trình từ thẻ SD qua SPI lên bộ nhớ chính trước khi thực thi.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Khảo sát nhu cầu ứng dụng thực tế cần bộ xử lý đa lõi",
    "- Khảo sát nhu cầu ứng dụng thực tế cần bộ xử lý đa lõi, các công trình liên quan, đặc tả RISC-V, AMBA AHB-Lite và AMBA AXI4; so sánh phạm vi chức năng, handshake, arbitration, burst, back-pressure và khả năng tích hợp DMA/DRAM để xác lập vai trò của từng bus trong kiến trúc đề xuất.",
)

replace_paragraph(
    content_cell.paragraphs,
    "Nội dung 3: Mở rộng thành hệ thống 4 lõi với interconnect AHB",
    "Nội dung 3: Mở rộng thành hệ thống 4 lõi với AHB-Lite interconnect, AXI4 system bus, cache L2 dùng chung và cơ chế coherence",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Mở rộng thiết kế lõi đơn thành hệ thống 4 lõi RV32IMA",
    "- Mở rộng thiết kế lõi đơn thành hệ thống 4 lõi RV32IMA, mỗi lõi có cache L1 và MMU/TLB riêng. Bốn lõi truy cập cache L2 dùng chung qua AHB-Lite interconnect; L2, DMA, DRAM và ngoại vi thuộc miền AXI4 system bus. Giao diện giữa hai miền bảo đảm chuyển đổi giao dịch, back-pressure và thứ tự truy cập nhất quán.",
)

replace_paragraph(
    content_cell.paragraphs,
    "- Thiết kế khối phân xử truy cập bộ nhớ dùng chung",
    "- Thiết kế arbiter/interconnect AHB-Lite nhiều master, giao diện AXI4 với các kênh đọc/ghi độc lập và burst transfer, bridge giữa hai miền bus, coherence manager và DMA; xây dựng testbench kiểm tra protocol assertions, timeout, back-pressure và tranh chấp truy cập.",
)

# Update the compact content-summary table heading.
summary = main.rows[6].cells[0].tables[0]
for row in summary.rows:
    for cell in row.cells:
        if cell.text.strip().startswith("Mở rộng thành hệ thống 4 lõi với interconnect AHB"):
            set_paragraph_text(
                cell.paragraphs[0],
                "Mở rộng thành hệ thống 4 lõi với AHB-Lite, AXI4, cache L2 dùng chung và coherence",
            )

# Update the detailed 20-week schedule. The same merged cell is exposed twice;
# editing the first instance updates the underlying XML once.
schedule = main.rows[6].cells[0].tables[1]
schedule_updates = {
    (2, 1): "Hoàn thiện testbench MMU/TLB/PTW: hit, miss, refill, page walk 2 cấp và page fault bằng page table mẫu.",
    (2, 2): "Chốt kiến trúc bus hai tầng; đặc tả AHB-Lite core/L1-L2 và AXI4 L2-DMA-DRAM/ngoại vi; lập ma trận địa chỉ và giao diện bridge.",
    (2, 3): "Đặc tả MMU/TLB, AHB-Lite và AXI4 được chốt ngay trong tháng 9.",
    (3, 1): "Hoàn thiện RTL và regression MMU/TLB/PTW; kiểm tra TLB flush, đồng thời truy cập lệnh/dữ liệu và lỗi dịch địa chỉ.",
    (3, 2): "Hiện thực AHB-Lite arbiter nhiều master và AXI4 read/write channel cơ bản; viết testbench handshake, wait-state và back-pressure.",
    (3, 3): "MMU/TLB PASS kiểm tra đơn vị; AHB-Lite và AXI4 hoàn thành luồng giao dịch cơ bản.",
    (4, 1): "Tích hợp MMU/TLB với lõi RV32IMA và cache L1; kiểm tra thứ tự VA - dịch địa chỉ - PA - truy cập cache/bus.",
    (4, 2): "Hoàn thiện AXI4 burst và bridge AHB-Lite/AXI4; bắt đầu DMA mem-to-mem trên miền AXI4.",
    (4, 3): "Mốc cuối tháng 9/đầu tháng 10: MMU/TLB, AHB-Lite, AXI4 và DMA v1 PASS ở mức đơn vị.",
    (5, 1): "Hoàn thiện khối lõi đơn RV32IMA + MMU/TLB + Cache L1 + AHB-Lite master; đối chiếu I/M/A với Spike.",
    (5, 2): "Tích hợp AHB-Lite interconnect với mô hình L2 và AXI4 system bus; kiểm tra chuyển tiếp giao dịch qua bridge.",
    (5, 3): "Khối lõi đơn và đường truy cập xuyên AHB-Lite/AXI4 PASS.",
    (6, 1): "Nhân bản khối lõi đơn đã có MMU/TLB và Cache L1 thành 4 lõi; đấu nối các AHB-Lite master.",
    (6, 2): "Hoàn thiện phân xử AHB-Lite 4 lõi, kiểm tra fairness/timeout; ổn định AXI4 system bus cho DMA và DRAM model.",
    (6, 3): "Mốc giữa tháng 10: bốn lõi truy cập qua AHB-Lite và đi đến miền AXI4 đúng giao thức.",
    (7, 1): "Hỗ trợ debug tích hợp MMU/TLB, cache L1 và giao tiếp AHB-Lite trên bốn lõi.",
    (7, 2): "Thiết kế và RTL cache L2 dùng chung; nối L2 với AHB-Lite phía lõi và AXI4 phía hệ thống.",
    (7, 3): "Cache L2 có thể nhận giao dịch từ bốn lõi và phát giao dịch xuống AXI4.",
    (8, 1): "Xây dựng regression/UVM ưu tiên cho MMU/TLB và AHB-Lite; khóa các lỗi giao tiếp trước khi mở rộng coherence.",
    (8, 2): "Tích hợp L2; bắt đầu coherence và bổ sung assertion/coverage cho AXI4 cùng đường bridge.",
    (8, 3): "Mốc cuối tháng 10: nền tảng MMU/TLB + AHB-Lite + AXI4 ổn định để triển khai coherence và DMA nâng cao.",
}
for (ri, ci), text in schedule_updates.items():
    set_paragraph_text(schedule.rows[ri].cells[ci].paragraphs[0], text)

# Keep the detailed schedule readable while reducing avoidable row splits.
for row in schedule.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9.5)

doc.save(OUTPUT)
print(OUTPUT)
