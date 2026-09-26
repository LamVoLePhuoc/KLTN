from docx import Document


def walk_table(table, path):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell_path = f"{path}/r{ri}c{ci}"
            for pi, paragraph in enumerate(cell.paragraphs):
                text = paragraph.text.strip()
                if text:
                    print(f"{cell_path}/p{pi} [{paragraph.style.name}]: {text}")
            for nti, nested in enumerate(cell.tables):
                walk_table(nested, f"{cell_path}/t{nti}")


doc = Document(r"C:\Users\ADMIN\Documents\GitHub\KLTN\23520840_23520838_DeCuongKLTN.docx")
for ti, table in enumerate(doc.tables):
    walk_table(table, f"t{ti}")
